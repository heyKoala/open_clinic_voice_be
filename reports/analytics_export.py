"""The Reports page as an Excel workbook: a summary sheet plus one sheet per table, for the same filters."""
from __future__ import annotations

from datetime import timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="0B1437")
TITLE_FONT = Font(bold=True, size=14)

LANGUAGES = {
    "en": "English", "hi": "Hindi", "kn": "Kannada", "ta": "Tamil", "te": "Telugu", "ml": "Malayalam",
    "mr": "Marathi", "bn": "Bengali", "gu": "Gujarati", "pa": "Punjabi", "ar": "Arabic",
}


def _table(sheet, headers, rows, widths):
    """A styled header row and its rows, appended after whatever the sheet already holds."""
    sheet.append(headers)
    start = sheet.max_row
    for cell in sheet[start]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(list(row))
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    return start


def _share(part, whole):
    return round(part / whole * 100, 1) if whole else 0


def _outcome_text(outcome: str) -> str:
    text = (outcome or "").replace("_", " ").replace("-", " ").strip()
    return text[:1].upper() + text[1:].lower() if text else ""


def _local(dt):
    # Excel has no time zones: write the clinic's local wall-clock time.
    return timezone.localtime(dt).replace(tzinfo=None)


def build_workbook(*, clinic, since, until, filters: dict, analytics: dict, doctors: list[dict], calls) -> bytes:
    s = analytics["summary"]
    first_day, last_day = timezone.localdate(since), timezone.localdate(until)
    wb = Workbook()

    # Summary
    summary = wb.active
    summary.title = "Summary"
    summary.append([f"AI call report: {clinic.name}"])
    summary["A1"].font = TITLE_FONT
    summary.append([f"{first_day:%d %b %Y} to {last_day:%d %b %Y}"])
    summary.append([f"Generated {_local(timezone.now()):%d %b %Y, %I:%M %p}"])
    summary.append([])
    for name, value in filters.items():
        summary.append([name, value])
    if s.get("test_calls"):
        summary.append(["Test calls", f"{s['test_calls']} left out of these numbers"])
    summary.append([])
    _table(summary, ["Metric", "Value"], [
        ("Total calls", s["total_calls"]),
        ("Inbound calls", s["inbound_calls"]),
        ("Outbound calls", s["outbound_calls"]),
        ("Appointments booked", s["calls_with_booking"]),
        ("Booked (% of calls)", s["booking_conversion_rate"]),
        ("Cancellations", s["calls_cancelled"]),
        ("Information only", s["calls_info_only"]),
        ("Unanswered", s["calls_unanswered"]),
        ("Other", s["calls_other"]),
        ("Average call length (seconds)", s["avg_duration_seconds"]),
        ("Total talk time (minutes)", round(s["total_duration_seconds"] / 60, 1)),
        ("Patients reached", s["patients_reached"]),
    ], widths=[32, 40])

    # Daily calls: every day of the period, zero where there were none.
    daily = wb.create_sheet("Daily calls")
    by_day = {row["date"]: row for row in analytics["calls_by_day"]}
    days = []
    day = first_day
    while day <= last_day:
        row = by_day.get(day.isoformat(), {"total": 0, "converted": 0})
        days.append((day, row["total"], row["converted"]))
        day += timedelta(days=1)
    _table(daily, ["Date", "Calls", "Appointments booked"], days, widths=[14, 10, 22])
    for cell in daily["A"][1:]:
        cell.number_format = "dd mmm yyyy"

    # Outcomes and languages
    total = s["total_calls"]
    outcomes = wb.create_sheet("Outcomes")
    _table(outcomes, ["Outcome", "Calls", "Share (%)"], [
        (name, count, _share(count, total)) for name, count in (
            ("Appointment booked", s["calls_with_booking"]), ("Cancelled", s["calls_cancelled"]),
            ("Information only", s["calls_info_only"]), ("Unanswered", s["calls_unanswered"]), ("Other", s["calls_other"]),
        )
    ], widths=[22, 10, 12])

    sentiment = wb.create_sheet("Sentiment")
    moods = analytics["sentiment_breakdown"]
    _table(sentiment, ["Caller sentiment", "Calls", "Share (%)"], [
        (label, moods.get(key, 0), _share(moods.get(key, 0), total))
        for key, label in (("positive", "Positive"), ("neutral", "Neutral"), ("negative", "Negative"), ("not_rated", "Not rated"))
    ], widths=[20, 10, 12])

    languages = wb.create_sheet("Languages")
    _table(languages, ["Language", "Calls", "Share (%)"], [
        (LANGUAGES.get(code, code), count, _share(count, total))
        for code, count in sorted(analytics["language_breakdown"].items(), key=lambda kv: -kv[1])
    ], widths=[16, 10, 12])

    # Doctors
    doctor_sheet = wb.create_sheet("Doctors")
    _table(doctor_sheet, ["Doctor", "Specialty", "Appointments", "Completed", "Completion (%)", "Avg. wait (min)", "Follow-up rate (%)"], [
        (f"Dr. {d['doctor_name']}", d.get("specialty") or "", d["total_appointments"], d["completed_appointments"],
         _share(d["completed_appointments"], d["total_appointments"]),
         round(d["avg_wait_seconds"] / 60, 1) if d["avg_wait_seconds"] else None, d["return_rate"])
        for d in doctors
    ], widths=[26, 20, 14, 12, 15, 15, 18])

    # Every call
    call_sheet = wb.create_sheet("Call logs")
    _table(call_sheet, ["Date & time", "Patient", "Phone", "Direction", "Language", "Duration (s)", "Outcome", "Sentiment", "Doctor", "Summary"], (
        (_local(c.occurred_at), c.patient.full_name if c.patient else "Unknown caller", c.patient.phone if c.patient else "",
         c.get_direction_display(), LANGUAGES.get(c.language, c.language), c.duration_seconds, _outcome_text(c.outcome),
         c.get_sentiment_display() if c.sentiment else "",
         f"Dr. {c.appointment.doctor.full_name}" if c.appointment_id and c.appointment.doctor_id else "", c.summary)
        for c in calls
    ), widths=[20, 24, 16, 11, 11, 13, 22, 12, 22, 70])
    for cell in call_sheet["A"][1:]:
        cell.number_format = "dd mmm yyyy hh:mm"
    for cell in call_sheet["J"][1:]:
        cell.alignment = Alignment(wrap_text=False)

    for sheet in wb.worksheets[1:]:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
