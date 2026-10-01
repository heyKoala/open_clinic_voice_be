"""Report generation service layer for ManageOPD exports."""

from __future__ import annotations

import csv
import os
import tempfile
from io import BytesIO

from django.utils import timezone
from django.utils.text import slugify
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


class ReportGenerator:
    """Generate report files for a report execution."""

    def __init__(self, report_execution):
        self.report_execution = report_execution
        self.template = report_execution.template
        self.parameters = report_execution.parameters
        self.clinic_id = report_execution.clinic_id

    def generate(self):
        """Generate the report and return the local file path."""
        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{slugify(self.template.name) or 'report'}_{timestamp}.{self.template.format}"

        if self.template.format == 'pdf':
            return self._generate_pdf(filename)
        elif self.template.format == 'csv':
            return self._generate_csv(filename)
        elif self.template.format == 'excel':
            return self._generate_excel(filename)
        else:
            return self._generate_text(filename)

    def _summary_rows(self):
        return [
            ("Report", self.template.name),
            ("Clinic ID", self.clinic_id),
            ("Report Type", self.template.get_report_type_display()),
            ("Format", self.template.get_format_display()),
            ("Generated", timezone.now().strftime('%Y-%m-%d %H:%M:%S UTC')),
            ("Retention (days)", self.template.retention_days),
            ("Allowed Roles", ", ".join(self.template.allowed_roles_display)),
            ("Parameters", str(self.parameters)),
        ]

    def _generate_pdf(self, filename):
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Heading1"],
            fontSize=24,
            spaceAfter=30
        )
        story.append(Paragraph(f"{self.template.name}", title_style))
        story.append(Spacer(1, 12))

        rows = self._summary_rows()
        table = Table([[key, str(value)] for key, value in rows], colWidths=[2.0 * inch, 4.8 * inch])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f8fafc")),
            ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#0f172a")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(table)
        story.append(Spacer(1, 18))
        story.append(Paragraph("Configuration", styles["Heading2"]))
        story.append(Paragraph(str(self.template.configuration or {}), styles["Code"] if "Code" in styles.byName else styles["Normal"]))

        doc.build(story)

        temp_path = os.path.join(tempfile.gettempdir(), filename)
        with open(temp_path, 'wb') as f:
            f.write(buffer.getvalue())

        return temp_path

    def _generate_csv(self, filename):
        """Generate a CSV report."""
        temp_path = os.path.join(tempfile.gettempdir(), filename)

        with open(temp_path, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(["Field", "Value"])
            for key, value in self._summary_rows():
                writer.writerow([key, value])

        return temp_path

    def _generate_excel(self, filename):
        wb = Workbook()
        summary = wb.active
        summary.title = "Summary"
        summary.append(["Field", "Value"])
        for cell in summary[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="0F172A")

        for key, value in self._summary_rows():
            summary.append([key, str(value)])

        summary.column_dimensions["A"].width = 24
        summary.column_dimensions["B"].width = 72
        for row in summary.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        config_sheet = wb.create_sheet("Configuration")
        config_sheet.append(["Key", "Value"])
        for cell in config_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="0F172A")
        for key, value in (self.template.configuration or {}).items():
            config_sheet.append([key, str(value)])
        config_sheet.column_dimensions["A"].width = 24
        config_sheet.column_dimensions["B"].width = 72

        temp_path = os.path.join(tempfile.gettempdir(), filename)
        wb.save(temp_path)

        return temp_path

    def _generate_text(self, filename):
        """Generate a text report."""
        temp_path = os.path.join(tempfile.gettempdir(), filename)

        with open(temp_path, 'w', encoding='utf-8') as f:
            f.write(f"{self.template.name}\n")
            f.write("=" * len(self.template.name) + "\n\n")
            for key, value in self._summary_rows():
                f.write(f"{key}: {value}\n")
            f.write("\nConfiguration:\n")
            for key, value in (self.template.configuration or {}).items():
                f.write(f"- {key}: {value}\n")

        return temp_path