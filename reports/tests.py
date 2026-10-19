from datetime import timedelta

import pytest
from rest_framework import status
from django.utils import timezone
from django.core.files.uploadedfile import SimpleUploadedFile

from accounts.models import User
from reports.models import ReportTemplate, ReportExecution
from appointments.models import Appointment
from followups.models import FollowUp
from queue_mgmt.models import QueueToken
from patients.models import Patient
from doctors.models import Doctor

pytestmark = pytest.mark.django_db


class TestReportingAPI:
    def test_doctor_sees_only_allowed_templates(self, doctor_client, admin_client, clinic):
        # Admin-only template
        ReportTemplate.objects.create(
            clinic=clinic,
            name="Revenue Report",
            report_type=ReportTemplate.ReportType.REVENUE_REPORT,
            format=ReportTemplate.ReportFormat.PDF,
            allowed_roles=[User.Role.CLINIC_ADMIN]
        )
        
        # Doctor template
        ReportTemplate.objects.create(
            clinic=clinic,
            name="Clinical Encounters",
            report_type=ReportTemplate.ReportType.CLINICAL_ENCOUNTERS,
            format=ReportTemplate.ReportFormat.PDF,
            allowed_roles=[User.Role.DOCTOR, User.Role.CLINIC_ADMIN]
        )

        res = doctor_client.get("/api/v1/reports/templates/")
        assert res.status_code == status.HTTP_200_OK
        
        # Determine if it's paginated or flat list
        data = res.data["results"] if "results" in res.data else res.data
        assert len(data) == 1
        assert data[0]["name"] == "Clinical Encounters"
        
        # Admin sees both
        res_admin = admin_client.get("/api/v1/reports/templates/")
        data_admin = res_admin.data["results"] if "results" in res_admin.data else res_admin.data
        assert len(data_admin) == 2

    def test_generate_from_template(self, doctor_client, doctor_user, clinic):
        from unittest.mock import patch
        with patch("reports.tasks.generate_report_task.delay") as mock_delay:
            template = ReportTemplate.objects.create(
                clinic=clinic,
                name="Clinical Encounters",
                report_type=ReportTemplate.ReportType.CLINICAL_ENCOUNTERS,
                format=ReportTemplate.ReportFormat.PDF,
                allowed_roles=[User.Role.DOCTOR]
            )
            
            res = doctor_client.post(
                "/api/v1/reports/executions/generate_from_template/",
                {"template_id": template.id, "parameters": {}},
                format="json"
            )
            assert res.status_code == status.HTTP_202_ACCEPTED
            assert res.data["status"] == "pending"
            
            mock_delay.assert_called_once()
        
    def test_signed_download_workflow(self, doctor_client, doctor_user, clinic):
        template = ReportTemplate.objects.create(
            clinic=clinic,
            name="Clinical Encounters",
            report_type=ReportTemplate.ReportType.CLINICAL_ENCOUNTERS,
            format=ReportTemplate.ReportFormat.PDF,
            allowed_roles=[User.Role.DOCTOR]
        )
        
        # Fake a completed report execution
        dummy_file = SimpleUploadedFile("test_report.pdf", b"file_content", content_type="application/pdf")
        execution = ReportExecution.objects.create(
            clinic=clinic,
            template=template,
            requested_by=doctor_user,
            status=ReportExecution.Status.COMPLETED,
            result_file=dummy_file,
            expires_at=timezone.now() + timezone.timedelta(days=7)
        )
        
        # 1. Get signed token
        res_sign = doctor_client.post(f"/api/v1/reports/executions/{execution.id}/signed-download/")
        assert res_sign.status_code == status.HTTP_200_OK
        download_url = res_sign.data["download_url"]
        
        # URL looks like: http://testserver/api/v1/reports/executions/<id>/download/?token=<token>
        token_qs = download_url.split("?token=")[-1]
        
        # 2. Download via GET using token
        res_dl = doctor_client.get(f"/api/v1/reports/executions/{execution.id}/download/?token={token_qs}")
        assert res_dl.status_code == status.HTTP_200_OK
        assert b"".join(res_dl.streaming_content) == b"file_content"
        
        # 3. Invalid token should fail
        res_fail = doctor_client.get(f"/api/v1/reports/executions/{execution.id}/download/?token=invalid_token")
        assert res_fail.status_code == status.HTTP_403_FORBIDDEN


class TestDashboardMetrics:
    def test_dashboard_metrics_computes_correctly(self, admin_client, clinic):
        # Create some data
        patient = Patient.objects.create(clinic=clinic, full_name="John Doe")
        doctor_user = User.objects.create_user(clinic=clinic, email="doc@example.com", password="pwd", role=User.Role.DOCTOR)
        doctor = Doctor.objects.create(clinic=clinic, user=doctor_user, full_name="Dr. Smith")
        
        now = timezone.now()
        # Appointment logic: 1 scheduled, 1 completed today
        Appointment.objects.create(clinic=clinic, patient=patient, doctor=doctor, starts_at=now, ends_at=now, status=Appointment.Status.SCHEDULED)
        Appointment.objects.create(clinic=clinic, patient=patient, doctor=doctor, starts_at=now, ends_at=now, status=Appointment.Status.COMPLETED)
        
        # Followup logic: 1 pending, 1 completed
        FollowUp.objects.create(clinic=clinic, patient=patient, doctor=doctor, scheduled_for=now, status=FollowUp.Status.PENDING)
        FollowUp.objects.create(clinic=clinic, patient=patient, doctor=doctor, scheduled_for=now, status=FollowUp.Status.COMPLETED)
        
        res = admin_client.get("/api/v1/reports/dashboard/")
        assert res.status_code == status.HTTP_200_OK
        
        data = res.data
        assert data["appointments"]["total"] == 2
        assert data["appointments"]["today"] == 2
        assert data["appointments"]["completion_rate"] == 50.0  # 1 completed out of 2 (scheduled+completed)
        
        assert data["followups"]["total"] == 2
        assert data["followups"]["due_today"] == 1  # only pending is due
        assert data["followups"]["conversion_rate"] == 100.0  # 1 completed out of 1 eligible (completed+failed+cancelled)

    def test_dashboard_isolates_clinic_data(self, admin_client, clinic, clinic_b):
        # Create data in Clinic B
        patient_b = Patient.objects.create(clinic=clinic_b, full_name="Jane Doe")
        doctor_user_b = User.objects.create_user(clinic=clinic_b, email="docb@example.com", password="pwd", role=User.Role.DOCTOR)
        doctor_b = Doctor.objects.create(clinic=clinic_b, user=doctor_user_b, full_name="Dr. B")
        
        now = timezone.now()
        Appointment.objects.create(clinic=clinic_b, patient=patient_b, doctor=doctor_b, starts_at=now, ends_at=now, status=Appointment.Status.SCHEDULED)
        
        # Admin in Clinic A checks their dashboard
        res = admin_client.get("/api/v1/reports/dashboard/")
        assert res.status_code == status.HTTP_200_OK
        assert res.data["appointments"]["total"] == 0

    def test_dashboard_restricted_to_admins(self, doctor_client, receptionist_client):
        assert doctor_client.get("/api/v1/reports/dashboard/").status_code == status.HTTP_403_FORBIDDEN
        assert receptionist_client.get("/api/v1/reports/dashboard/").status_code == status.HTTP_403_FORBIDDEN


class TestAIAnalytics:
    def _call(self, clinic, when, **fields):
        from ai_agent.models import CallLog
        defaults = {"direction": "inbound", "language": "en", "outcome": "Appointment booked", "duration_seconds": 60}
        return CallLog.objects.create(clinic=clinic, occurred_at=when, **{**defaults, **fields})

    def test_date_range_and_filters(self, admin_client, clinic):
        patient = Patient.objects.create(clinic=clinic, full_name="John Doe")
        now = timezone.now()
        self._call(clinic, now, patient=patient)
        self._call(clinic, now, patient=patient, language="hi", outcome="Cancelled", duration_seconds=30)
        self._call(clinic, now, direction="outbound", outcome="Info only")
        self._call(clinic, now - timedelta(days=40))  # outside the range

        today = timezone.localdate().isoformat()
        res = admin_client.get(f"/api/v1/reports/ai-analytics/?start={today}&end={today}")
        assert res.status_code == status.HTTP_200_OK
        summary = res.data["summary"]
        assert summary["total_calls"] == 3
        assert summary["calls_with_booking"] == 1
        assert summary["calls_cancelled"] == 1
        assert summary["patients_reached"] == 1
        assert summary["total_duration_seconds"] == 150
        assert res.data["language_breakdown"] == {"en": 2, "hi": 1}
        assert res.data["period_days"] == 1

        res = admin_client.get(f"/api/v1/reports/ai-analytics/?start={today}&direction=inbound&language=en")
        assert res.data["summary"]["total_calls"] == 1

        # Without a range it still covers the last ?days.
        assert admin_client.get("/api/v1/reports/ai-analytics/?days=60").data["summary"]["total_calls"] == 4

    def test_call_logs_filter_by_outcome_keyword(self, admin_client, clinic):
        now = timezone.now()
        booked = self._call(clinic, now)
        self._call(clinic, now, outcome="Cancelled")
        res = admin_client.get("/api/v1/ai/call-logs/?outcome=booked")
        assert [row["id"] for row in res.data["results"]] == [booked.id]

    def test_front_desk_sees_call_analytics_but_doctors_do_not(self, receptionist_client, doctor_client):
        assert receptionist_client.get("/api/v1/reports/ai-analytics/").status_code == status.HTTP_200_OK
        assert doctor_client.get("/api/v1/reports/ai-analytics/").status_code == status.HTTP_403_FORBIDDEN
        # The clinic-wide and per-doctor dashboards stay admin-only.
        assert receptionist_client.get("/api/v1/reports/doctor-analytics/").status_code == status.HTTP_403_FORBIDDEN


class TestLiveQueue:
    def test_live_queue_computes_wait_times(self, doctor_client, clinic, doctor_user):
        patient = Patient.objects.create(clinic=clinic, full_name="John Doe")
        doctor = Doctor.objects.create(clinic=clinic, user=doctor_user, full_name="Dr. Smith")
        
        today = timezone.localdate()
        now = timezone.now()
        check_in_time = now - timezone.timedelta(minutes=15)
        
        # 1 token waiting for 15 minutes
        QueueToken.objects.create(
            clinic=clinic,
            patient=patient,
            doctor=doctor,
            service_date=today,
            token_number=1,
            status=QueueToken.Status.WAITING,
            checked_in_at=check_in_time
        )
        
        res = doctor_client.get("/api/v1/reports/live-queue/")
        assert res.status_code == status.HTTP_200_OK
        
        summary = res.data["summary"]
        assert summary["total"] == 1
        assert summary["waiting"] == 1
        assert summary["average_wait_seconds"] >= 900  # at least 15 minutes

class TestDoctorFilterAndExport:
    def _setup(self, clinic):
        from ai_agent.models import CallLog
        patient = Patient.objects.create(clinic=clinic, full_name="Asha Menon", phone="9876543210")
        doctors = []
        for name in ("Anvi Pillai", "Darsh Sanghvi"):
            user = User.objects.create_user(clinic=clinic, email=f"{name.split()[0].lower()}@example.com", password="pwd", role=User.Role.DOCTOR)
            doctors.append(Doctor.objects.create(clinic=clinic, user=user, full_name=name))
        now = timezone.now()
        appts = [Appointment.objects.create(clinic=clinic, patient=patient, doctor=d, starts_at=now, ends_at=now) for d in doctors]
        CallLog.objects.create(clinic=clinic, patient=patient, direction="inbound", occurred_at=now, outcome="appointment_booked", appointment=appts[0], duration_seconds=60)
        CallLog.objects.create(clinic=clinic, patient=patient, direction="inbound", occurred_at=now, outcome="appointment_cancelled", appointment=appts[1], duration_seconds=30)
        CallLog.objects.create(clinic=clinic, direction="inbound", occurred_at=now, outcome="information_only")
        return doctors

    def test_doctor_filter(self, admin_client, clinic):
        anvi, darsh = self._setup(clinic)
        today = timezone.localdate().isoformat()
        everyone = admin_client.get(f"/api/v1/reports/ai-analytics/?start={today}").data["summary"]
        anvis = admin_client.get(f"/api/v1/reports/ai-analytics/?start={today}&doctor={anvi.id}").data["summary"]
        assert (everyone["total_calls"], anvis["total_calls"], anvis["calls_with_booking"], anvis["calls_cancelled"]) == (3, 1, 1, 0)

        rows = admin_client.get(f"/api/v1/ai/call-logs/?doctor={darsh.id}").data["results"]
        assert [(r["outcome"], r["doctor"], r["doctor_name"]) for r in rows] == [("appointment_cancelled", darsh.id, "Darsh Sanghvi")]

        performance = admin_client.get(f"/api/v1/reports/doctor-analytics/?start={today}&doctor={anvi.id}").data["doctors"]
        assert [d["doctor_id"] for d in performance] == [anvi.id]

    def test_excel_download_follows_the_filters(self, admin_client, receptionist_client, clinic):
        from io import BytesIO
        from openpyxl import load_workbook
        anvi, _ = self._setup(clinic)
        today = timezone.localdate().isoformat()

        res = admin_client.get(f"/api/v1/reports/ai-analytics/export/?start={today}&doctor={anvi.id}")
        assert res.status_code == status.HTTP_200_OK
        assert res["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
        assert f'call-report_{today}_to_{today}.xlsx' in res["Content-Disposition"]

        wb = load_workbook(BytesIO(res.content))
        assert wb.sheetnames == ["Summary", "Daily calls", "Outcomes", "Sentiment", "Languages", "Doctors", "Call logs"]
        summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0]}
        assert summary["Doctor"] == "Dr. Anvi Pillai" and summary["Total calls"] == 1
        calls = list(wb["Call logs"].iter_rows(min_row=2, values_only=True))
        assert [(c[1], c[6], c[8]) for c in calls] == [("Asha Menon", "Appointment booked", "Dr. Anvi Pillai")]
        assert [r[0] for r in wb["Doctors"].iter_rows(min_row=2, values_only=True)] == ["Dr. Anvi Pillai"]

        assert receptionist_client.get("/api/v1/reports/ai-analytics/export/").status_code == status.HTTP_403_FORBIDDEN
